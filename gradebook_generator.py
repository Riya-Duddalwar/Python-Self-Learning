from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUTPUT_FILE = "AIML_Batch_Grades.xlsx"
SHEET_TITLE = "Semester_Results"
PASS_MARK = 40

HEADERS = [
    "Roll No",
    "Student Name",
    "Lab Marks (out of 50)",
    "Theory Marks (out of 50)",
    "Total Marks (out of 100)",
    "Result",
]

STUDENTS = [
    (101, "Aarav Sharma", 42, 38),
    (102, "Diya Patel", 35, 41),
    (103, "Rohan Mehta", 18, 15),
    (104, "Sneha Iyer", 47, 44),
    (105, "Kabir Singh", 28, 9),
    (106, "Ananya Desai", 39, 36),
]


def create_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    return workbook, sheet


def add_headers(sheet):
    sheet.append(HEADERS)


def add_students(sheet):
    for roll_no, name, lab, theory in STUDENTS:
        sheet.append([roll_no, name, lab, theory])


def calculate_results(sheet):
    totals = []
    for row in range(2, sheet.max_row + 1):
        lab = sheet.cell(row=row, column=3).value
        theory = sheet.cell(row=row, column=4).value
        total = lab + theory
        result = "Pass" if total >= PASS_MARK else "Fail"
        sheet.cell(row=row, column=5, value=total)
        sheet.cell(row=row, column=6, value=result)
        totals.append(total)
    return totals


def add_class_average(sheet, totals):
    average = round(sum(totals) / len(totals), 2)
    summary_row = sheet.max_row + 1
    sheet.cell(row=summary_row, column=2, value="Class Average")
    sheet.cell(row=summary_row, column=5, value=average)
    sheet.cell(row=summary_row, column=2).font = Font(bold=True)
    sheet.cell(row=summary_row, column=5).font = Font(bold=True)
    return average


def style_headers(sheet):
    gray_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = gray_fill


def adjust_column_widths(sheet):
    for column in sheet.columns:
        longest = max(len(str(cell.value)) for cell in column if cell.value is not None)
        sheet.column_dimensions[column[0].column_letter].width = longest + 4


def main():
    workbook, sheet = create_sheet()
    add_headers(sheet)
    add_students(sheet)
    totals = calculate_results(sheet)
    average = add_class_average(sheet, totals)
    style_headers(sheet)
    adjust_column_widths(sheet)
    workbook.save(OUTPUT_FILE)
    print("Saved", OUTPUT_FILE)
    print("Class average:", average)


if __name__ == "__main__":
    main()
